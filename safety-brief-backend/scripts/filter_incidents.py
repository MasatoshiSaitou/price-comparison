"""職場のあんぜんサイト 労働災害（死傷）データベースから、
「製造業」かつ「試験・実験・開発・測定」に関連する事例を抽出するスクリプト。

このスクリプトは、厚生労働省のサイト (anzeninfo.mhlw.go.jp) からブラウザで
ダウンロードした CSV / Excel ファイルを入力にとる。ネットワークアクセスは
行わない（ダウンロードはあらかじめ手動で行っておくこと）。

pandas は使わない（コンパイル済みバイナリが企業PCのアプリケーション制御
ポリシーでブロックされる環境があったため）。CSV は標準ライブラリの csv、
Excel は純Python実装の openpyxl のみを使う。

実行方法:
    pip install openpyxl
    python filter_incidents.py "data/*.xlsx"

出力: manufacturing_test_incidents.csv （抽出された事例の一覧）
"""

import csv
import glob
import sys

# 職場のあんぜんサイトの労働災害（死傷）データベースの列順（22列固定）
COLUMNS = [
    "id",
    "era",
    "year_in_era",
    "month",
    "time_range",
    "description",
    "industry_major_code",
    "industry_major_name",
    "industry_middle_code",
    "industry_middle_name",
    "industry_minor_code",
    "industry_minor_name",
    "workplace_size",
    "cause_major_code",
    "cause_major_name",
    "cause_middle_code",
    "cause_middle_name",
    "cause_minor_code",
    "cause_minor_name",
    "accident_type_code",
    "accident_type_name",
    "age",
]

TARGET_INDUSTRY = "製造業"
KEYWORDS = ["試験", "実験", "開発", "測定"]

HEADER_ROWS_TO_SKIP = 2


def load_csv_rows(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        rows = list(reader)
    return rows[HEADER_ROWS_TO_SKIP:]


def load_xlsx_rows(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl がインストールされていません。`pip install openpyxl` を実行してください。"
        ) from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < HEADER_ROWS_TO_SKIP:
            continue
        rows.append(["" if v is None else str(v) for v in row])
    wb.close()
    return rows


def load_rows(path):
    if path.lower().endswith(".csv"):
        return load_csv_rows(path)
    return load_xlsx_rows(path)


def to_records(rows):
    records = []
    for row in rows:
        if len(row) < len(COLUMNS):
            continue
        records.append(dict(zip(COLUMNS, row)))
    return records


def main(paths):
    if not paths:
        print("使い方: python filter_incidents.py <ファイルパスまたはglobパターン> ...")
        return

    expanded_paths = []
    for p in paths:
        matched = glob.glob(p)
        expanded_paths.extend(matched if matched else [p])

    all_records = []
    for path in expanded_paths:
        try:
            rows = load_rows(path)
            records = to_records(rows)
            all_records.extend(records)
            print(f"読み込み成功: {path} ({len(records)}件)")
        except Exception as exc:  # noqa: BLE001
            print(f"スキップ: {path} ({exc})")

    if not all_records:
        print("読み込めるファイルがありませんでした。")
        return

    manufacturing = [r for r in all_records if r.get("industry_major_name") == TARGET_INDUSTRY]
    matched = [
        r
        for r in manufacturing
        if any(keyword in (r.get("description") or "") for keyword in KEYWORDS)
    ]

    print(f"\n読み込んだ全件数: {len(all_records)}")
    print(f"製造業: {len(manufacturing)}件")
    print(f"うちキーワード（{'/'.join(KEYWORDS)}）一致: {len(matched)}件")

    output_path = "manufacturing_test_incidents.csv"
    with open(output_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(matched)
    print(f"-> {output_path} に出力しました")


if __name__ == "__main__":
    main(sys.argv[1:])
